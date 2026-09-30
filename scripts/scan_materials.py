"""Read local source materials into a private audit file; never publish the output."""

from pathlib import Path
import json

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / ".private-audit" / "extracted.json"
OUT.parent.mkdir(exist_ok=True)
records = []

for path in sorted(ROOT.rglob("*")):
    if not path.is_file() or any(part in {".git", ".next", "node_modules", ".private-audit", "public"} for part in path.relative_to(ROOT).parts):
        continue
    if path.suffix.lower() not in {".pdf", ".docx", ".jpg", ".jpeg", ".png"}:
        continue
    record = {"path": str(path.relative_to(ROOT)), "bytes": path.stat().st_size, "mtime": path.stat().st_mtime}
    try:
        if path.suffix.lower() == ".pdf":
            reader = PdfReader(path)
            record["pages"] = len(reader.pages)
            record["text"] = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]
        elif path.suffix.lower() == ".docx":
            doc = Document(path)
            record["text"] = [p.text for p in doc.paragraphs if p.text.strip()]
            record["tables"] = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
        else:
            from PIL import Image

            with Image.open(path) as image:
                record["image"] = {"width": image.width, "height": image.height}
    except Exception as exc:
        record["error"] = str(exc)
    records.append(record)

OUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Scanned {len(records)} files; private audit written to {OUT}")
for record in records:
    text = record.get("text", [])
    chars = sum(len(page) for page in text)
    print(record["path"], f"pages={record.get('pages', '-')}", f"chars={chars}", f"error={record.get('error', '-')}")
